import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / ".docx_qa"))

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = ROOT / "docs" / "EDUCONNECT_THIET_KE_VAI_TRO.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, value=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = text
        cell.width = Cm(widths[idx])
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row_idx, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for col_idx, text in enumerate(values):
            cells[col_idx].text = text
            cells[col_idx].width = Cm(widths[col_idx])
            set_cell_margins(cells[col_idx])
            cells[col_idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx % 2 == 0:
                set_cell_shading(cells[col_idx], "EAF2F8")
            for p in cells[col_idx].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(2)
section.right_margin = Cm(2)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.5)
styles["Title"].font.name = "Arial"
styles["Title"].font.size = Pt(23)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
for name, size in (("Heading 1", 16), ("Heading 2", 12.5)):
    styles[name].font.name = "Arial"
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

doc.add_heading("Thiết kế vai trò và chức năng EduConnect", 0)
p = doc.add_paragraph("Đề xuất phân quyền ngắn gọn cho phiên bản tiếp theo")
p.runs[0].italic = True
p.runs[0].font.color.rgb = RGBColor(89, 89, 89)

doc.add_paragraph(
    "EduConnect nên có năm nhóm người dùng: Khách, Tài khoản gia đình, Giáo viên, "
    "Quản lý hệ thống và Quản trị viên. Tài khoản gia đình dùng chung thông tin đăng nhập "
    "nhưng tách hồ sơ phụ huynh và từng học sinh để lịch học, tiến độ và quyền riêng tư không bị lẫn."
)

doc.add_heading("Nguyên tắc phân quyền", level=1)
add_bullets(doc, [
    "Mỗi người chỉ thấy dữ liệu và thao tác cần cho vai trò của mình.",
    "Phụ huynh là chủ tài khoản gia đình; mỗi học sinh có hồ sơ riêng bên trong tài khoản.",
    "Quản lý hệ thống vận hành hằng ngày; Quản trị viên quản lý quyền, bảo mật và cấu hình.",
    "Mọi thao tác nhạy cảm phải được ghi nhật ký để kiểm tra lại.",
])

doc.add_heading("Tổng quan các vai trò", level=1)
add_table(doc, ["Vai trò", "Mục đích", "Phạm vi chính"], [
    ["Khách", "Tìm hiểu trước khi đăng ký", "Xem thông tin công khai, tìm giáo viên, xem tài liệu mẫu và đăng ký"],
    ["Tài khoản gia đình", "Quản lý việc học của học sinh", "Đặt lịch, thanh toán, theo dõi tiến độ, học tập và trao đổi"],
    ["Giáo viên", "Cung cấp hoạt động dạy học", "Hồ sơ chuyên môn, lịch dạy, lớp học, tài liệu, tiến độ và thu nhập"],
    ["Quản lý hệ thống", "Vận hành nền tảng hằng ngày", "Duyệt nội dung, hỗ trợ người dùng, xử lý sự cố và xem báo cáo"],
    ["Quản trị viên", "Kiểm soát toàn bộ hệ thống", "Phân quyền, bảo mật, cấu hình, nhật ký, tích hợp và dữ liệu"],
], [3.1, 5.0, 8.2])

doc.add_heading("Chức năng theo từng vai trò", level=1)

role_sections = [
    ("Khách chưa đăng ký", [
        "Xem giới thiệu, dịch vụ, bảng giá, câu hỏi thường gặp và thông tin liên hệ.",
        "Tìm kiếm, lọc và xem hồ sơ công khai của giáo viên.",
        "Xem trước tài liệu miễn phí và đánh giá công khai.",
        "Đăng ký, đăng nhập hoặc gửi yêu cầu tư vấn.",
        "Không được đặt lịch, nhắn tin hoặc xem dữ liệu cá nhân.",
    ]),
    ("Tài khoản gia đình gồm phụ huynh và học sinh", [
        "Dùng một tài khoản đăng nhập; tạo nhiều hồ sơ học sinh trong gia đình.",
        "Phụ huynh quản lý hồ sơ, lịch học, đặt lịch, thanh toán, hoàn tiền và đưa đón.",
        "Phụ huynh xem tiến độ, nhận xét, điểm số, bài tập và thông báo của từng học sinh.",
        "Học sinh xem lịch, vào buổi học, nhận tài liệu, nộp bài và xem tiến độ của mình.",
        "Phụ huynh quyết định học sinh có được tự nhắn tin, đặt lịch hoặc xem chi phí hay không.",
        "Gia đình đánh giá giáo viên, lưu giáo viên yêu thích và gửi yêu cầu hỗ trợ.",
    ]),
    ("Giáo viên", [
        "Đăng ký hồ sơ, xác minh danh tính, bằng cấp và chứng chỉ.",
        "Quản lý môn dạy, học phí, hình thức dạy, khu vực và lịch rảnh.",
        "Nhận, xác nhận, đổi lịch hoặc từ chối yêu cầu đặt học.",
        "Quản lý học sinh, điểm danh, nội dung buổi học và báo cáo tiến độ.",
        "Giao bài, nhận bài nộp, chia sẻ tài liệu và phản hồi.",
        "Nhắn tin với gia đình và xử lý yêu cầu liên quan đến lớp học.",
        "Xem doanh thu, phí nền tảng và yêu cầu rút tiền.",
    ]),
    ("Quản lý hệ thống", [
        "Duyệt hoặc từ chối hồ sơ giáo viên, chứng chỉ và tài liệu.",
        "Xác minh, cảnh báo, tạm khóa và mở khóa người dùng theo quy trình.",
        "Theo dõi đặt lịch, thanh toán, hoàn tiền, khiếu nại và hỗ trợ.",
        "Quản lý môn học, cấp học, khu vực, nội dung trang và thông báo chung.",
        "Xem báo cáo truy cập, tìm kiếm, lượt nhấp, đặt lịch, doanh thu và chất lượng.",
        "Không được tự cấp quyền Quản trị viên hoặc thay đổi cấu hình bảo mật quan trọng.",
    ]),
    ("Quản trị viên", [
        "Tạo và thu hồi tài khoản Quản lý hệ thống; gán hoặc thay đổi vai trò.",
        "Thiết lập quyền, xác thực hai bước, thời hạn đăng nhập và giới hạn truy cập.",
        "Quản lý cấu hình, thanh toán, email, lưu trữ và các dịch vụ tích hợp.",
        "Xem nhật ký hoạt động, cảnh báo bảo mật, sao lưu và tình trạng hệ thống.",
        "Khóa, khôi phục hoặc xóa tài khoản trong trường hợp đặc biệt.",
        "Mọi thao tác nhạy cảm vẫn phải lưu nhật ký và yêu cầu xác nhận lại.",
    ]),
]
for heading, bullets in role_sections:
    doc.add_heading(heading, level=2)
    add_bullets(doc, bullets)

doc.add_heading("Phân quyền trong tài khoản gia đình", level=1)
add_table(doc, ["Chức năng", "Phụ huynh", "Học sinh"], [
    ["Quản lý hồ sơ gia đình và thêm học sinh", "Có", "Không"],
    ["Đặt hoặc thay đổi lịch học", "Có", "Theo cho phép"],
    ["Thanh toán, hoàn tiền và xem chi phí", "Có", "Không"],
    ["Xem lịch và tham gia buổi học", "Có", "Có"],
    ["Xem tiến độ, bài tập và tài liệu", "Có", "Chỉ hồ sơ của mình"],
    ["Nhắn tin với giáo viên", "Có", "Theo cho phép"],
    ["Đánh giá giáo viên", "Có", "Có thể góp ý"],
], [8.2, 3.8, 4.3])

doc.add_heading("Chức năng nên bổ sung", level=1)
add_table(doc, ["Ưu tiên", "Chức năng", "Lý do"], [
    ["Cao", "Hồ sơ gia đình và hồ sơ học sinh", "Dùng chung tài khoản nhưng dữ liệu học tập vẫn tách rõ"],
    ["Cao", "Phân quyền và nhật ký thao tác", "Ngăn truy cập sai và biết ai đã thay đổi dữ liệu"],
    ["Cao", "Thông báo và nhắn tin", "Giảm bỏ lỡ lịch học, bài tập và phản hồi"],
    ["Cao", "Thanh toán, hoàn tiền và thu nhập", "Hoàn chỉnh quy trình đặt học"],
    ["Trung bình", "Bài tập và phòng học trực tuyến", "Hỗ trợ toàn bộ quá trình học"],
    ["Trung bình", "Khiếu nại, báo xấu và hỗ trợ", "Xử lý tranh chấp và bảo vệ người dùng"],
    ["Trung bình", "Xuất hoặc xóa dữ liệu cá nhân", "Tăng quyền kiểm soát dữ liệu"],
    ["Sau", "Mã giảm giá, gói học, giới thiệu bạn bè", "Bổ sung sau khi luồng chính ổn định"],
], [2.5, 6.0, 7.8])

doc.add_heading("Thứ tự triển khai đề xuất", level=1)
add_bullets(doc, [
    "Giai đoạn 1: Chuẩn hóa năm vai trò, quyền truy cập, hồ sơ gia đình và nhật ký thao tác.",
    "Giai đoạn 2: Hoàn thiện đặt lịch, thanh toán, thông báo, nhắn tin và quy trình dạy học.",
    "Giai đoạn 3: Bổ sung hỗ trợ, khiếu nại, báo cáo nâng cao và tính năng tăng trưởng.",
])

doc.add_heading("Kết luận", level=1)
doc.add_paragraph(
    "Mô hình này giữ trải nghiệm đơn giản cho gia đình nhưng vẫn tách dữ liệu của từng học sinh. "
    "Quản lý hệ thống phụ trách vận hành hằng ngày; Quản trị viên chịu trách nhiệm về quyền hạn, "
    "bảo mật và cấu hình toàn nền tảng."
)

footer = section.footer.paragraphs[0]
footer.text = "EduConnect | Thiết kế vai trò và chức năng"
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in footer.runs:
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(117, 117, 117)

doc.save(OUT)
print(OUT)
